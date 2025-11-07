"""
Simple Excel Export - Bare Minimum
Just read simple_qto table and dump to Excel with basic formatting
"""

import sqlite3
import sys
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError:
    print("ERROR: openpyxl not installed")
    print("Install with: pip install openpyxl")
    sys.exit(1)


def export_to_excel(db_path: str, output_path: str = None):
    """
    Export simple_qto table to Excel
    """

    if output_path is None:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = f"QTO_Report_{timestamp}.xlsx"

    print("\n" + "="*80)
    print("EXCEL EXPORT")
    print("="*80 + "\n")

    # Read data from database
    print(f"Reading data from: {db_path}")
    conn = sqlite3.connect(db_path)

    # Check if simple_qto table exists
    cursor = conn.execute("""
        SELECT COUNT(*) FROM sqlite_master
        WHERE type='table' AND name='simple_qto'
    """)

    if cursor.fetchone()[0] == 0:
        print("\nERROR: simple_qto table not found!")
        print("Run this first: python simple_qto_extract.py <database_path>")
        conn.close()
        sys.exit(1)

    # Get all data
    cursor = conn.execute("""
        SELECT
            discipline,
            ifc_class,
            measurement_type,
            element_count,
            total_quantity,
            uom,
            avg_quantity
        FROM simple_qto
        ORDER BY measurement_type, total_quantity DESC
    """)

    rows = cursor.fetchall()
    conn.close()

    if not rows:
        print("ERROR: No data in simple_qto table")
        sys.exit(1)

    print(f"Found {len(rows)} QTO line items\n")

    # Create Excel workbook
    print(f"Creating Excel file: {output_path}")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "QTO Summary"

    # Title
    ws['A1'] = "BIM QUANTITY TAKE-OFF"
    ws['A1'].font = Font(size=16, bold=True)
    ws.merge_cells('A1:G1')

    ws['A2'] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws['A2'].font = Font(size=9, italic=True)

    ws['A3'] = f"Database: {db_path}"
    ws['A3'].font = Font(size=9, italic=True)

    # Headers
    headers = ['Discipline', 'IFC Class', 'Type', 'Count', 'Total Qty', 'UOM', 'Avg Qty']
    row = 5

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    # Data rows
    row += 1
    for data_row in rows:
        for col, value in enumerate(data_row, start=1):
            cell = ws.cell(row, col, value)

            # Format numbers
            if col in [4, 5, 7]:  # Count, Total Qty, Avg Qty columns
                if isinstance(value, (int, float)):
                    cell.number_format = '#,##0.00'

        row += 1

    # Auto-size columns
    for col_idx in range(1, 8):  # 7 columns
        max_length = 0
        column_letter = openpyxl.utils.get_column_letter(col_idx)

        for row_idx in range(1, row + 1):
            cell = ws.cell(row_idx, col_idx)
            if cell.value and not isinstance(cell, openpyxl.cell.cell.MergedCell):
                max_length = max(max_length, len(str(cell.value)))

        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    # Freeze header row
    ws.freeze_panes = 'A6'

    # Add summary sheet
    ws_summary = wb.create_sheet("Summary by Discipline")

    # Get summary data
    conn = sqlite3.connect(db_path)
    cursor = conn.execute("""
        SELECT
            discipline,
            measurement_type,
            SUM(element_count) AS total_elements,
            ROUND(SUM(total_quantity), 2) AS total_qty,
            uom
        FROM simple_qto
        GROUP BY discipline, measurement_type, uom
        ORDER BY discipline, measurement_type
    """)

    summary_rows = cursor.fetchall()
    conn.close()

    # Summary title
    ws_summary['A1'] = "QTO SUMMARY BY DISCIPLINE"
    ws_summary['A1'].font = Font(size=14, bold=True)
    ws_summary.merge_cells('A1:E1')

    # Summary headers
    summary_headers = ['Discipline', 'Type', 'Elements', 'Total Quantity', 'UOM']
    row = 3

    for col, header in enumerate(summary_headers, start=1):
        cell = ws_summary.cell(row, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    # Summary data
    row += 1
    for data_row in summary_rows:
        for col, value in enumerate(data_row, start=1):
            cell = ws_summary.cell(row, col, value)

            if col in [3, 4]:  # Elements, Total Quantity
                if isinstance(value, (int, float)):
                    cell.number_format = '#,##0.00'

        row += 1

    # Auto-size summary columns
    for col_idx in range(1, 6):  # 5 columns
        max_length = 0
        column_letter = openpyxl.utils.get_column_letter(col_idx)

        for row_idx in range(1, row + 1):
            cell = ws_summary.cell(row_idx, col_idx)
            if cell.value and not isinstance(cell, openpyxl.cell.cell.MergedCell):
                max_length = max(max_length, len(str(cell.value)))

        adjusted_width = min(max_length + 2, 50)
        ws_summary.column_dimensions[column_letter].width = adjusted_width

    ws_summary.freeze_panes = 'A4'

    # Save workbook
    wb.save(output_path)

    print("\n" + "="*80)
    print(f"✓ EXCEL FILE CREATED: {output_path}")
    print("="*80 + "\n")

    print("Sheets created:")
    print("  1. QTO Summary - All elements with quantities")
    print("  2. Summary by Discipline - Rolled up totals")
    print()

    return output_path


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("Usage: python simple_excel_export.py <database_path> [output_path]")
        print("Example: python simple_excel_export.py ~/Documents/bonsai/DatabaseFiles/sample_extracted_v3.db")
        sys.exit(1)

    db_path = sys.argv[1]
    output_path = sys.argv[2] if len(sys.argv) > 2 else None

    export_to_excel(db_path, output_path)
